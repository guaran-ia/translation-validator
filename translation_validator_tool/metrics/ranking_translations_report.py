#!/usr/bin/env python3
"""
Genera un reporte Word con el ranking de traducciones por usuario
en un rango de fechas personalizado.
"""

import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path

import psycopg2
from psycopg2 import Error


DATABASE_URL = "postgresql://postgres:1234@localhost:5432/translation_validator"
SERVER_AHEAD_HOURS = 3
TIME_SUB_SQL = f" - INTERVAL '{SERVER_AHEAD_HOURS} hours'"
REPORT_FILENAME = "Ranking_Traducciones_Por_Fecha.docx"
TOP_LIMIT = 20


def get_docx_tools():
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except ImportError:
        print(
            "Falta la dependencia python-docx. "
            "Instalala con: pip install -r requirements_metrics.txt"
        )
        sys.exit(1)

    return Document, WD_ALIGN_PARAGRAPH, Pt, RGBColor


def parse_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        raise argparse.ArgumentTypeError(
            f"Fecha invalida: {value}. Usa el formato YYYY-MM-DD."
        )


def get_connection():
    try:
        return psycopg2.connect(DATABASE_URL)
    except Error as e:
        print(f"Error al conectar con la base de datos: {e}")
        sys.exit(1)


def ensure_documents_folder():
    documents_folder = Path(__file__).parent / "documents"
    documents_folder.mkdir(exist_ok=True)
    return documents_folder


def get_translation_ranking(conn, start_date, end_date):
    query = f"""
    SELECT
        u.id AS user_id,
        u.nombre_completo,
        u.email,
        COUNT(*) AS traducciones_total
    FROM annotation_translations at
    INNER JOIN users u
        ON u.id = at.user_id
    WHERE (at.created_at{TIME_SUB_SQL}) >= %s
      AND (at.created_at{TIME_SUB_SQL}) < %s
    GROUP BY
        u.id,
        u.nombre_completo,
        u.email
    ORDER BY
        COUNT(*) DESC,
        u.nombre_completo ASC
    LIMIT %s;
    """

    end_exclusive = end_date + timedelta(days=1)

    try:
        with conn.cursor() as cursor:
            cursor.execute(query, (start_date, end_exclusive, TOP_LIMIT))
            return cursor.fetchall()
    except Error as e:
        print(f"Error ejecutando query de ranking: {e}")
        return None


def add_title(doc, start_date, end_date, text_tools):
    WD_ALIGN_PARAGRAPH, Pt, RGBColor = text_tools

    title = doc.add_paragraph()
    title_run = title.add_run(
        f"Reporte desde: {start_date.isoformat()} hasta: {end_date.isoformat()}"
    )
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0, 102, 204)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run(
        f"Top {TOP_LIMIT} usuarios con mas traducciones realizadas"
    )
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    generated = doc.add_paragraph()
    generated_run = generated.add_run(
        f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    generated_run.font.size = Pt(9)
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("_" * 80)


def add_ranking_table(doc, rows):
    if not rows:
        doc.add_paragraph("No se encontraron traducciones en el rango seleccionado.")
        return

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["Posicion", "Nombre completo", "Email", "Total de traducciones"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for position, row in enumerate(rows, start=1):
        user_id, full_name, email, translations_total = row
        cells = table.add_row().cells
        cells[0].text = str(position)
        cells[1].text = full_name or "Usuario sin nombre"
        cells[2].text = email or "-"
        cells[3].text = str(translations_total)


def generate_docx(rows, start_date, end_date, documents_folder):
    Document, WD_ALIGN_PARAGRAPH, Pt, RGBColor = get_docx_tools()
    doc = Document()
    add_title(doc, start_date, end_date, (WD_ALIGN_PARAGRAPH, Pt, RGBColor))
    add_ranking_table(doc, rows)

    output_path = documents_folder / REPORT_FILENAME
    doc.save(str(output_path))
    return output_path


def build_parser():
    parser = argparse.ArgumentParser(
        description="Genera un Word con el top 20 de traducciones por usuario."
    )
    parser.add_argument(
        "--desde",
        required=True,
        type=parse_date,
        help="Fecha inicial inclusiva en formato YYYY-MM-DD.",
    )
    parser.add_argument(
        "--hasta",
        required=True,
        type=parse_date,
        help="Fecha final inclusiva en formato YYYY-MM-DD.",
    )
    return parser


def main():
    parser = build_parser()
    args = parser.parse_args()

    if args.desde > args.hasta:
        parser.error("--desde no puede ser mayor que --hasta")

    conn = get_connection()

    try:
        rows = get_translation_ranking(conn, args.desde, args.hasta)
        if rows is None:
            sys.exit(1)

        documents_folder = ensure_documents_folder()
        output_path = generate_docx(rows, args.desde, args.hasta, documents_folder)

        print(f"Reporte generado en: {output_path}")
        print(f"Usuarios incluidos: {len(rows)}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
