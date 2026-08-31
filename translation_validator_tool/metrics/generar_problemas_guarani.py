#!/usr/bin/env python3
"""Exporta 10 problemas matemáticos validados y traducidos al guaraní."""

import argparse
import json
import os
import sys
from collections import Counter
from pathlib import Path

import psycopg2
from openpyxl import load_workbook
from psycopg2 import Error
from psycopg2.extras import RealDictCursor


LIMIT = 10
MIN_QUALITY = 4.0
DEFAULT_DATABASE_URL = "postgresql://postgres:1234@localhost:5432/translation_validator"

SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = SCRIPT_DIR / "documents" / "problemas_matematicos_guarani.json"
CORPUS_FILE = SCRIPT_DIR.parent.parent / "data" / "data.xlsx"
CORPUS_SHEETS = ("test_mgsm_es_to_gn", "train_mgsm_es_to_gn")


CANDIDATES_QUERY = """
WITH validation_summary AS (
    SELECT
        av.translation_id,
        COUNT(*) FILTER (
            WHERE av.status = 'completed'
              AND av.suggestion_id IS NULL
              AND av.part_type IN ('question', 'both')
        )::int AS completed_count,
        COUNT(*) FILTER (
            WHERE av.status = 'completed'
              AND av.suggestion_id IS NULL
              AND av.part_type IN ('question', 'both')
              AND av.fluency_score IS NOT NULL
              AND av.accuracy_score IS NOT NULL
        )::int AS scored_count,
        AVG((av.fluency_score + av.accuracy_score) / 2.0) FILTER (
            WHERE av.status = 'completed'
              AND av.suggestion_id IS NULL
              AND av.part_type IN ('question', 'both')
              AND av.fluency_score IS NOT NULL
              AND av.accuracy_score IS NOT NULL
        ) AS quality_average,
        BOOL_AND(av.correctness = 'correct') FILTER (
            WHERE av.status = 'completed'
              AND av.suggestion_id IS NULL
              AND av.part_type IN ('question', 'both')
              AND av.fluency_score IS NOT NULL
              AND av.accuracy_score IS NOT NULL
        ) AS all_correct,
        MAX(av.updated_at) FILTER (
            WHERE av.status = 'completed'
              AND av.suggestion_id IS NULL
              AND av.part_type IN ('question', 'both')
              AND av.fluency_score IS NOT NULL
              AND av.accuracy_score IS NOT NULL
        ) AS latest_validation_at
    FROM annotation_validations av
    GROUP BY av.translation_id
)
SELECT
    at.id AS translation_id,
    a.id AS annotation_id,
    a.question_es,
    at.question_gn,
    a.answer_number,
    a.md5,
    at.validated,
    COALESCE(vs.completed_count, 0) AS completed_count,
    COALESCE(vs.scored_count, 0) AS scored_count,
    vs.quality_average,
    COALESCE(vs.all_correct, false) AS all_correct,
    vs.latest_validation_at
FROM annotation_translations at
INNER JOIN annotations a
    ON a.id = at.annotation_id
LEFT JOIN validation_summary vs
    ON vs.translation_id = at.id
ORDER BY at.id ASC;
"""


def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Genera documents/problemas_matematicos_guarani.json "
            "con exactamente 10 problemas validados."
        )
    )
    return parser.parse_args()


def normalize_number(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def load_corpus_answers():
    if not CORPUS_FILE.exists():
        print(
            f"Advertencia: no se encontró el corpus para respuestas de respaldo: {CORPUS_FILE}",
            file=sys.stderr,
        )
        return {}

    answers = {}
    workbook = load_workbook(CORPUS_FILE, read_only=True, data_only=True)
    try:
        for sheet_name in CORPUS_SHEETS:
            if sheet_name not in workbook.sheetnames:
                continue

            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue

            columns = {name: index for index, name in enumerate(headers) if name}
            if "question_md5" not in columns or "answer_number" not in columns:
                continue

            md5_index = columns["question_md5"]
            answer_index = columns["answer_number"]
            for row in rows:
                if len(row) <= max(md5_index, answer_index):
                    continue
                question_md5 = row[md5_index]
                answer_number = row[answer_index]
                if question_md5 and answer_number is not None:
                    answers[str(question_md5).strip()] = normalize_number(answer_number)
    finally:
        workbook.close()

    return answers


def fetch_candidates(database_url):
    connection = None
    try:
        connection = psycopg2.connect(database_url)
        connection.set_session(readonly=True, autocommit=True)
        with connection.cursor(cursor_factory=RealDictCursor) as cursor:
            cursor.execute(CANDIDATES_QUERY)
            return cursor.fetchall()
    finally:
        if connection is not None:
            connection.close()


def evaluate_candidates(rows, corpus_answers):
    discarded = Counter()
    eligible = []

    for row in rows:
        reasons = []
        problem = (row["question_gn"] or "").strip()
        problem_es = (row["question_es"] or "").strip()
        expected_answer = row["answer_number"]
        if expected_answer is None and row["md5"]:
            expected_answer = corpus_answers.get(row["md5"].strip())

        if not row["validated"]:
            reasons.append("traducción no marcada como validada")
        if not problem:
            reasons.append("traducción al guaraní vacía")
        if not problem_es:
            reasons.append("premisa original en español vacía")
        if row["completed_count"] == 0:
            reasons.append("sin validaciones completadas")
        elif row["scored_count"] != row["completed_count"]:
            reasons.append("validaciones completadas sin puntuaciones válidas")
        if row["scored_count"] > 0 and not row["all_correct"]:
            reasons.append("alguna validación no está calificada como correcta")

        quality_average = row["quality_average"]
        if quality_average is None:
            reasons.append("sin promedio de calidad calculable")
        elif float(quality_average) < MIN_QUALITY:
            reasons.append("promedio de calidad inferior a 4")
        if expected_answer is None:
            reasons.append("sin respuesta matemática esperada")

        if reasons:
            discarded.update(reasons)
            continue

        eligible.append(
            {
                "annotation_id": row["annotation_id"],
                "problem": problem,
                "problema_es": problem_es,
                "expected_answer": normalize_number(expected_answer),
                "quality_average": float(quality_average),
                "validation_count": row["scored_count"],
                "latest_validation_at": row["latest_validation_at"],
            }
        )

    eligible.sort(
        key=lambda item: (
            item["quality_average"],
            item["validation_count"],
            item["latest_validation_at"],
        ),
        reverse=True,
    )

    selected = []
    seen_annotations = set()
    for candidate in eligible:
        if candidate["annotation_id"] in seen_annotations:
            discarded["premisa original duplicada"] += 1
            continue
        seen_annotations.add(candidate["annotation_id"])
        selected.append(candidate)
        if len(selected) == LIMIT:
            break

    return selected, discarded


def validate_and_serialize(selected):
    if len(selected) != LIMIT:
        raise ValueError(f"se esperaban {LIMIT} registros y se obtuvieron {len(selected)}")

    annotation_ids = [item["annotation_id"] for item in selected]
    if len(annotation_ids) != len(set(annotation_ids)):
        raise ValueError("existen premisas originales duplicadas")

    problems = []
    for index, item in enumerate(selected, start=1):
        if not item["problem"]:
            raise ValueError(f"el problema {index} está vacío")
        if not item["problema_es"]:
            raise ValueError(f"el problema {index} no tiene problema_es")
        if item["expected_answer"] is None:
            raise ValueError(f"el problema {index} no tiene expected_answer")
        if item["quality_average"] < MIN_QUALITY:
            raise ValueError(f"el problema {index} tiene promedio inferior a {MIN_QUALITY}")

        problems.append(
            {
                "id": index,
                "problem": item["problem"],
                "problema_es": item["problema_es"],
                "expected_answer": item["expected_answer"],
                "difficulty": "fácil",
            }
        )

    content = json.dumps(problems, ensure_ascii=False, indent=2) + "\n"
    parsed = json.loads(content)
    if not isinstance(parsed, list) or len(parsed) != LIMIT:
        raise ValueError("el JSON final no tiene la estructura esperada")
    return content


def write_output(content):
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary_file = OUTPUT_FILE.with_suffix(OUTPUT_FILE.suffix + ".tmp")
    try:
        temporary_file.write_text(content, encoding="utf-8")
        temporary_file.replace(OUTPUT_FILE)
    finally:
        if temporary_file.exists():
            temporary_file.unlink()


def print_discard_summary(discarded):
    print("Causas de descarte (una traducción puede tener más de una causa):")
    if not discarded:
        print("  - Ninguna")
        return
    for reason, count in discarded.most_common():
        print(f"  - {reason}: {count}")


def main():
    parse_args()
    database_url = os.getenv("DATABASE_URL", DEFAULT_DATABASE_URL)

    try:
        corpus_answers = load_corpus_answers()
        rows = fetch_candidates(database_url)
        selected, discarded = evaluate_candidates(rows, corpus_answers)

        if len(selected) < LIMIT:
            print(
                f"No se generó el archivo: se encontraron {len(selected)} "
                f"traducciones válidas de las {LIMIT} requeridas.",
                file=sys.stderr,
            )
            print_discard_summary(discarded)
            return 1

        content = validate_and_serialize(selected)
        write_output(content)

        averages = [item["quality_average"] for item in selected]
        print(f"Archivo generado: {OUTPUT_FILE}")
        print(f"Problemas exportados: {len(selected)}")
        print(f"Promedio mínimo: {min(averages):.2f}")
        print(f"Promedio máximo: {max(averages):.2f}")
        return 0
    except Error as error:
        print(f"Error de base de datos: {error}", file=sys.stderr)
        return 1
    except (OSError, ValueError, KeyError) as error:
        print(f"Error al generar el JSON: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
