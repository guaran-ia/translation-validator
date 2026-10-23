#!/usr/bin/env python3
"""
Script para crear usuarios masivamente desde archivo Excel.
Uso: 
  python bulk_create_users.py <archivo.xlsx>
  python bulk_create_users.py users.xlsx
  python bulk_create_users.py users.xlsx --role ambos  # Por defecto es "ambos"
  python bulk_create_users.py users.xlsx --dry-run  # Simula sin insertar en BD
"""
import sys
import os
import argparse
import logging
from pathlib import Path
from typing import List, Dict, Tuple

# Add API directory to Python path
current_dir = Path(__file__).resolve().parent
project_root = current_dir.parent
sys.path.insert(0, str(project_root / "api"))

# Set working directory to API for .env file loading
original_cwd = os.getcwd()
os.chdir(str(project_root / "api"))

try:
    import openpyxl
except ImportError:
    print("❌ openpyxl no está instalado. Instálalo con: pip install openpyxl")
    sys.exit(1)

try:
    from sqlalchemy import text
    from app.database import SessionLocal, engine
    from app.models.user import User
    from app.core.security import get_password_hash
    from app.config import settings
except ImportError as e:
    print(f"❌ Error importando módulos de la API: {e}")
    print("Asegúrate de estar en el directorio correcto")
    sys.exit(1)

logger = logging.getLogger(__name__)


def setup_logging(verbose: bool = False):
    """Configurar logging."""
    level = logging.DEBUG if verbose else logging.INFO
    logging.basicConfig(
        level=level,
        format='%(asctime)s - %(levelname)s - %(message)s'
    )
    return logging.getLogger(__name__)


def validate_excel_file(filepath: str) -> Path:
    """Validar que el archivo Excel existe."""
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"❌ Archivo no encontrado: {filepath}")
    if path.suffix.lower() not in ['.xlsx', '.xls']:
        raise ValueError(f"❌ El archivo debe ser .xlsx o .xls: {filepath}")
    return path


def read_excel_data(filepath: Path) -> List[Dict[str, str]]:
    """
    Leer datos del Excel.
    Retorna lista de dicts con: {nombre, edad, email}
    """
    try:
        logger.info(f"📂 Leyendo archivo: {filepath}")
        wb = openpyxl.load_workbook(filepath, data_only=True)
        
        # Buscar hoja "Base de Datos"
        if "Base de Datos" not in wb.sheetnames:
            logger.error(f"❌ Hoja 'Base de Datos' no encontrada. Hojas disponibles: {wb.sheetnames}")
            raise ValueError(f"Hoja 'Base de Datos' no encontrada")
        
        ws = wb["Base de Datos"]
        logger.info("✅ Hoja 'Base de Datos' encontrada")
        
        users_data = []
        max_row = ws.max_row
        logger.info(f"📋 Procesando filas 3 a {max_row}")
        
        # Iterar desde fila 3
        for row_idx in range(3, max_row + 1):
            cell_a = ws[f'A{row_idx}'].value  # Nombre
            cell_b = ws[f'B{row_idx}'].value  # Edad
            cell_c = ws[f'C{row_idx}'].value  # Email
            
            # Saltar filas vacías
            if not all([cell_a, cell_b, cell_c]):
                if cell_a or cell_b or cell_c:
                    logger.warning(f"⚠️  Fila {row_idx} incompleta: saltando...")
                continue
            
            # Convertir edad a string si es número
            edad = str(int(cell_b)) if isinstance(cell_b, (int, float)) else str(cell_b).strip()
            
            users_data.append({
                "nombre": str(cell_a).strip(),
                "edad": edad,
                "email": str(cell_c).strip().lower(),  # Normalizar a minúsculas
            })
        
        logger.info(f"✅ Se leyeron {len(users_data)} usuarios válidos")
        return users_data
        
    except Exception as e:
        logger.error(f"❌ Error leyendo Excel: {e}")
        raise


def process_user_data(user_data: Dict[str, str]) -> Dict[str, str]:
    """
    Procesar datos del usuario:
    - username: parte antes del @ del email
    - password: primera palabra del nombre + edad
    - email: tal cual
    """
    email = user_data['email']
    nombre = user_data['nombre']
    edad = user_data['edad']
    
    # Username: parte antes del @
    username = email.split('@')[0].lower()
    
    # Password: primera palabra del nombre + edad
    primera_palabra = nombre.split()[0]  # Primera palabra antes del espacio
    password = f"{primera_palabra}{edad}"
    
    return {
        "username": username,
        "email": email,
        "password": password,
        "original_name": nombre,
    }


def check_duplicates(db_session, processed_users: List[Dict]) -> Tuple[List[Dict], List[Dict]]:
    """
    Verificar duplicados en la BD y en la lista.
    Retorna: (usuarios_validos, usuarios_duplicados)
    """
    valid_users = []
    duplicate_users = []
    seen_usernames = set()
    seen_emails = set()
    
    for user in processed_users:
        # Verificar si ya existe en BD
        existing_user = db_session.query(User).filter(
            (User.username == user['username']) | (User.email == user['email'])
        ).first()
        
        if existing_user:
            logger.warning(
                f"⚠️  Duplicado en BD: {user['username']} ({user['email']}) - "
                f"user ya existe con id={existing_user.id}"
            )
            duplicate_users.append(user)
            continue
        
        # Verificar duplicados dentro de la lista actual
        if user['username'] in seen_usernames:
            logger.warning(f"⚠️  Username duplicado en lista: {user['username']}")
            duplicate_users.append(user)
            continue
        
        if user['email'] in seen_emails:
            logger.warning(f"⚠️  Email duplicado en lista: {user['email']}")
            duplicate_users.append(user)
            continue
        
        seen_usernames.add(user['username'])
        seen_emails.add(user['email'])
        valid_users.append(user)
    
    return valid_users, duplicate_users


def create_users_in_db(db_session, valid_users: List[Dict], role: str = "ambos") -> Tuple[int, List[str]]:
    """
    Crear usuarios en la base de datos.
    Retorna: (cantidad_creada, lista_de_errores)
    """
    created_count = 0
    errors = []
    
    for user_data in valid_users:
        try:
            hashed_password = get_password_hash(user_data['password'])
            
            new_user = User(
                username=user_data['username'],
                email=user_data['email'],
                hashed_password=hashed_password,
                role=role,
                is_active=True,
                is_superuser=False,
            )
            
            db_session.add(new_user)
            db_session.flush()  # Para obtener el ID
            
            logger.info(
                f"✅ Usuario creado: {user_data['username']} ({user_data['email']}) "
                f"ID={new_user.id} | Password={user_data['password']}"
            )
            created_count += 1
            
        except Exception as e:
            error_msg = f"❌ Error creando usuario {user_data['username']}: {str(e)}"
            logger.error(error_msg)
            errors.append(error_msg)
    
    if created_count > 0:
        try:
            db_session.commit()
            logger.info(f"✅ {created_count} usuarios guardados en la BD")
        except Exception as e:
            logger.error(f"❌ Error guardando usuarios en BD: {e}")
            db_session.rollback()
            errors.append(f"Error en commit: {str(e)}")
            created_count = 0
    
    return created_count, errors


def main():
    parser = argparse.ArgumentParser(
        description="Crear usuarios masivamente desde Excel"
    )
    parser.add_argument("excel_file", help="Archivo Excel con usuarios (users.xlsx)")
    parser.add_argument("--role", default="ambos", 
                       choices=["admin", "ambos"],
                       help="Role de los usuarios (default: ambos)")
    parser.add_argument("--dry-run", action="store_true",
                       help="Simular sin insertar en BD")
    parser.add_argument("-v", "--verbose", action="store_true",
                       help="Modo verbose")
    
    args = parser.parse_args()
    
    logger = setup_logging(args.verbose)
    
    print("=" * 70)
    print("🚀 BULK CREATE USERS - Translation Validator")
    print("=" * 70)
    
    try:
        # 1. Validar archivo Excel
        logger.info("📋 Validando archivo Excel...")
        excel_path = validate_excel_file(args.excel_file)
        
        # 2. Leer datos
        logger.info("📖 Leyendo datos del Excel...")
        users_raw = read_excel_data(excel_path)
        
        if not users_raw:
            logger.error("❌ No se encontraron usuarios en el Excel")
            sys.exit(1)
        
        # 3. Procesar datos
        logger.info("⚙️  Procesando datos...")
        processed_users = [process_user_data(u) for u in users_raw]
        
        # 4. Conectar a BD
        logger.info("🔌 Conectando a la base de datos...")
        db = SessionLocal()
        
        # 5. Verificar duplicados
        logger.info("🔍 Verificando duplicados...")
        valid_users, duplicate_users = check_duplicates(db, processed_users)
        
        logger.info(f"📊 Resumen:")
        logger.info(f"   - Total en archivo: {len(processed_users)}")
        logger.info(f"   - Válidos para crear: {len(valid_users)}")
        logger.info(f"   - Duplicados encontrados: {len(duplicate_users)}")
        
        if not valid_users:
            logger.warning("⚠️  No hay usuarios válidos para crear")
            return
        
        # 6. Crear usuarios
        if args.dry_run:
            logger.info("🧪 DRY-RUN MODE - No se crearán usuarios")
            print("\n" + "=" * 70)
            print("Usuarios que se crearían:")
            print("=" * 70)
            for user in valid_users:
                print(f"  Username: {user['username']}")
                print(f"  Email:    {user['email']}")
                print(f"  Password: {user['password']}")
                print(f"  Role:     {args.role}")
                print("-" * 70)
        else:
            logger.info(f"💾 Creando {len(valid_users)} usuarios en la BD...")
            created_count, errors = create_users_in_db(db, valid_users, role=args.role)
            
            print("\n" + "=" * 70)
            print(f"✅ RESULTADO: {created_count}/{len(valid_users)} usuarios creados")
            print("=" * 70)
            
            if errors:
                print("\n⚠️  Errores encontrados:")
                for error in errors:
                    print(f"  - {error}")
            
            if created_count == len(valid_users):
                print("\n🎉 ¡Todos los usuarios creados exitosamente!")
            elif created_count > 0:
                print(f"\n⚠️  Creados parcialmente: {created_count} de {len(valid_users)}")
            else:
                print("\n❌ No se pudo crear ningún usuario")
        
        db.close()
        
    except Exception as e:
        logger.error(f"❌ Error fatal: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
    finally:
        print("=" * 70)


if __name__ == "__main__":
    main()
