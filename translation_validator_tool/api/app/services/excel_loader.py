"""
Service for loading annotations from Excel files.
"""
from typing import List, Dict, Optional, Tuple
from pathlib import Path
import io

try:
    import openpyxl
except ImportError:
    openpyxl = None

from sqlalchemy.orm import Session
from app.models.annotation import Annotation


def get_excel_sheets(file_bytes: bytes) -> List[str]:
    """
    Get list of sheet names from an Excel file.
    
    Args:
        file_bytes: Excel file as bytes
        
    Returns:
        List of sheet names
        
    Raises:
        ValueError: If file is not a valid Excel file
    """
    if not openpyxl:
        raise ImportError("openpyxl is required for Excel support")
    
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        return workbook.sheetnames
    except Exception as e:
        raise ValueError(f"Invalid Excel file: {str(e)}")


def get_excel_preview(file_bytes: bytes, sheet_name: str) -> Dict[str, any]:
    """
    Get columns and first data row from an Excel sheet for preview/mapping.
    
    Args:
        file_bytes: Excel file as bytes
        sheet_name: Name of the sheet to preview
        
    Returns:
        Dictionary with:
        - columns: List of column names (from first row)
        - preview_row: Dictionary of first data row (row 2) with column names as keys
    """
    if not openpyxl:
        raise ImportError("openpyxl is required for Excel support")
    
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found")
        
        sheet = workbook[sheet_name]
        
        # Get headers from first row
        headers = []
        for col_idx in range(1, sheet.max_column + 1):
            header = sheet.cell(row=1, column=col_idx).value
            if header:
                headers.append(str(header).strip())
        
        # Get first data row (row 2)
        preview_row = {}
        for col_idx, header in enumerate(headers, 1):
            value = sheet.cell(row=2, column=col_idx).value
            preview_row[header] = str(value).strip() if value else ""
        
        return {
            "columns": headers,
            "preview_row": preview_row
        }
    except Exception as e:
        raise ValueError(f"Error reading Excel preview: {str(e)}")


def load_annotations_from_excel_bytes(
    file_bytes: bytes,
    sheet_name: str,
    db: Session,
    column_mapping: Optional[Dict[str, str]] = None,
    skip_duplicates: bool = True
) -> Dict[str, any]:
    """
    Load annotations from Excel file bytes using column mapping.
    
    Args:
        file_bytes: Excel file as bytes
        sheet_name: Name of the sheet to process
        db: Database session
        column_mapping: Dictionary mapping field names to column names:
                       {
                           "question_es": "column_name",
                           "answer_es": "column_name",
                           "answer_number": "column_name" (optional),
                           "equation_solution": "column_name" (optional)
                       }
        skip_duplicates: Whether to skip duplicate entries
        
    Returns:
        Dictionary with:
        - added_count: Number of annotations successfully added
        - skipped_count: Number of rows skipped
        - errors: List of error messages
        - details: List of added annotation details
    """
    if not openpyxl:
        raise ImportError("openpyxl is required for Excel support")
    
    result = {
        "added_count": 0,
        "skipped_count": 0,
        "errors": [],
        "details": []
    }
    
    # Use default hardcoded mapping if not provided (backward compatibility)
    if not column_mapping:
        column_mapping = {
            "question_es": "B",
            "answer_es": "D",
            "answer_number": "F",
            "equation_solution": "G"
        }
    
    try:
        # Load workbook
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as e:
            result["errors"].append(f"Error loading Excel file: {str(e)}")
            return result
        
        # Check sheet exists
        if sheet_name not in workbook.sheetnames:
            result["errors"].append(
                f"Sheet '{sheet_name}' not found. Available sheets: {', '.join(workbook.sheetnames)}"
            )
            return result
        
        sheet = workbook[sheet_name]
        
        # Build column index map from header row
        headers = {}
        for col_idx in range(1, sheet.max_column + 1):
            header_val = sheet.cell(row=1, column=col_idx).value
            if header_val:
                headers[str(header_val).strip()] = col_idx
        
        # Get column indices for each field
        col_indices = {}
        required_fields = {"question_es", "answer_es"}
        
        for field, col_name in column_mapping.items():
            if col_name:
                # Handle both column letter format (A, B, C) and column name format
                if len(col_name) == 1 and col_name.isalpha():
                    # Column letter format (A=1, B=2, etc.) - only if single character
                    col_indices[field] = ord(col_name.upper()) - ord('A') + 1
                elif col_name in headers:
                    # Column name format
                    col_indices[field] = headers[col_name]
                else:
                    # Column not found
                    if field in required_fields:
                        # Only report error for required fields
                        result["errors"].append(f"Column '{col_name}' not found in sheet")
                    # For optional fields, silently skip them if not found
        
        # Validate required fields have mappings
        if not col_indices.get("question_es") and not col_indices.get("answer_es"):
            result["errors"].append("At least one of 'question_es' or 'answer_es' must be mapped")
            return result
        
        # Process data rows starting from row 2
        row_num = 2
        
        while True:
            # Read values from mapped columns
            question_es = None
            answer_es = None
            answer_number = None
            equation_solution = None
            
            if "question_es" in col_indices:
                val = sheet.cell(row=row_num, column=col_indices["question_es"]).value
                question_es = str(val).strip() if val else None
            
            if "answer_es" in col_indices:
                val = sheet.cell(row=row_num, column=col_indices["answer_es"]).value
                answer_es = str(val).strip() if val else None
            
            # Check if row is empty
            if not question_es and not answer_es:
                break
            
            if "answer_number" in col_indices:
                val = sheet.cell(row=row_num, column=col_indices["answer_number"]).value
                if val:
                    try:
                        answer_number = int(val)
                    except (ValueError, TypeError):
                        result["errors"].append(
                            f"Row {row_num}: answer_number '{val}' is not a valid integer"
                        )
            
            if "equation_solution" in col_indices:
                val = sheet.cell(row=row_num, column=col_indices["equation_solution"]).value
                equation_solution = str(val).strip() if val else None
            
            # Check for duplicates if skip_duplicates is True
            if skip_duplicates:
                query = db.query(Annotation)
                
                if question_es and answer_es:
                    exists = query.filter(
                        Annotation.question_es == question_es,
                        Annotation.answer_es == answer_es
                    ).first()
                elif question_es:
                    exists = query.filter(Annotation.question_es == question_es).first()
                else:
                    exists = query.filter(Annotation.answer_es == answer_es).first()
                
                if exists:
                    result["skipped_count"] += 1
                    row_num += 1
                    continue
            
            # Create annotation record
            try:
                annotation = Annotation(
                    question_es=question_es,
                    answer_es=answer_es,
                    answer_number=answer_number,
                    equation_solution=equation_solution,
                    source="GuaranIA"
                )
                db.add(annotation)
                db.commit()
                db.refresh(annotation)
                
                result["added_count"] += 1
                result["details"].append({
                    "id": annotation.id,
                    "question_es": question_es[:60] + "..." if question_es and len(question_es) > 60 else question_es,
                    "answer_es": answer_es[:60] + "..." if answer_es and len(answer_es) > 60 else answer_es,
                    "equation_solution": equation_solution
                })
                
            except Exception as e:
                result["errors"].append(f"Row {row_num}: {str(e)}")
                result["skipped_count"] += 1
                db.rollback()
            
            row_num += 1
        
        return result
        
    except Exception as e:
        result["errors"].append(f"Error processing file: {str(e)}")
        db.rollback()
        return result
